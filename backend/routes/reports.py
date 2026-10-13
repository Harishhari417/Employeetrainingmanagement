import csv
import io
from datetime import datetime
from flask import Blueprint, jsonify, Response, request, send_file
from flask_jwt_extended import jwt_required, get_jwt
from bson import ObjectId
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from db import db

reports_bp = Blueprint("reports", __name__, url_prefix="/api/reports")

def allowed_training(claims, training):
    if not training: return False
    if claims.get("role") == "HR_ADMIN": return True
    if claims.get("role") == "MANAGER":
        if claims.get("department") in training.get("departments", []): return True
        ids={x.get("employeeId") for x in db.employees.find({"department":claims.get("department")},{"employeeId":1,"_id":0})}
        return db.training_participants.count_documents({"trainingId":str(training["_id"]),"employeeId":{"$in":list(ids)}})>0
    return training.get("trainerEmployeeId") == claims.get("employeeId") or db.training_participants.count_documents({"trainingId":str(training["_id"]),"employeeId":claims.get("employeeId")}) > 0

def training_or_404(training_id):
    if not ObjectId.is_valid(training_id): return None
    return db.trainings.find_one({"_id": ObjectId(training_id)})

def csv_response(rows, filename):
    out=io.StringIO(); fields=list(rows[0].keys()) if rows else ["No data"]
    w=csv.DictWriter(out, fieldnames=fields); w.writeheader(); w.writerows(rows)
    return Response(out.getvalue(), mimetype="text/csv", headers={"Content-Disposition": f"attachment; filename={filename}"})

def xlsx_response(sheets, filename):
    wb=Workbook(); first=True
    for title, rows in sheets:
        ws=wb.active if first else wb.create_sheet(); first=False; ws.title=title[:31]
        if not rows:
            ws.append(["No data"]); continue
        headers=list(rows[0].keys()); ws.append(headers)
        for cell in ws[1]:
            cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="3730A3"); cell.alignment=Alignment(vertical="center")
        for row in rows:
            values=[]
            for header in headers:
                value=row.get(header, "")
                if hasattr(value, "isoformat"):
                    try: value=value.replace(tzinfo=None).isoformat(sep=" ")
                    except TypeError: value=value.isoformat()
                elif value is None: value=""
                values.append(value)
            ws.append(values)
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            width=min(max(max(len(str(c.value or "")) for c in col)+2,12),42)
            ws.column_dimensions[get_column_letter(col[0].column)].width=width
    output=io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def attendance_rows(training_id):
    rows=[]
    for p in db.training_participants.find({"trainingId":training_id}):
        e=db.employees.find_one({"employeeId":p.get("employeeId")},{"name":1,"department":1,"_id":0}) or {}
        rows.append({"Employee ID":p.get("employeeId",""),"Employee":e.get("name",""),"Department":e.get("department",""),"Attendance":p.get("attendance","Pending"),"Marked By":p.get("attendanceMarkedBy",""),"Marked At":p.get("attendanceMarkedAt","")})
    return rows

def feedback_rows(training_id):
    rows=[]
    for f in db.feedback.find({"trainingId":training_id}).sort("submittedAt",-1):
        e=db.employees.find_one({"employeeId":f.get("employeeId")},{"name":1,"department":1,"_id":0}) or {}
        ratings=f.get("ratings",{})
        rows.append({"Employee ID":f.get("employeeId",""),"Employee":e.get("name",""),"Department":e.get("department",""),"Programme Design":ratings.get("Programme Design",""),"Presentation":ratings.get("Presentation of Information",""),"Information Amount":ratings.get("Amount of Information",""),"Administration":ratings.get("Administrative arrangements",""),"Overall":ratings.get("overall",""),"Major Learning":f.get("majorLearning",""),"Remarks":f.get("remarks",""),"Status":f.get("status","Pending"),"Submitted At":f.get("submittedAt","")})
    return rows

def trainee_trainer_rows(training_id, training):
    rows=[]
    for p in db.training_participants.find({"trainingId":training_id}):
        e=db.employees.find_one({"employeeId":p.get("employeeId")},{"name":1,"department":1,"designation":1,"_id":0}) or {}
        rows.append({"Trainee ID":p.get("employeeId",""),"Trainee":e.get("name",""),"Department":e.get("department",""),"Designation":e.get("designation",""),"Trainer":training.get("trainerName",""),"Trainer Employee ID":training.get("trainerEmployeeId",""),"Attendance":p.get("attendance","Pending"),"Feedback":p.get("feedbackStatus","Pending"),"Completion":p.get("completionStatus","Not Started")})
    return rows

@reports_bp.get("/trainings")
@jwt_required()
def report_trainings():
    claims=get_jwt(); query={}
    if claims.get("role")=="MANAGER":
        ids=[x.get("employeeId") for x in db.employees.find({"department":claims.get("department")},{"employeeId":1,"_id":0}) if x.get("employeeId")]
        query["$or"]=[{"departments":claims.get("department")},{"traineeIds":{"$in":ids}}]
    elif claims.get("role")=="EMPLOYEE": query["$or"]=[{"trainerEmployeeId":claims.get("employeeId")},{"traineeIds":claims.get("employeeId")}]
    rows=[]
    for t in db.trainings.find(query).sort("startDate",-1):
        tid=str(t["_id"]); count=db.training_participants.count_documents({"trainingId":tid})
        rows.append({"_id":tid,"title":t.get("title",""),"startDate":t.get("startDate",t.get("trainingDate")),"endDate":t.get("endDate",t.get("startDate",t.get("trainingDate"))),"trainerName":t.get("trainerName",""),"trainingType":t.get("trainingType",""),"status":t.get("status","Upcoming"),"trainees":count})
    return jsonify(rows)

@reports_bp.get("/training/<training_id>/combined")
@jwt_required()
def combined_training_report(training_id):
    claims=get_jwt(); t=training_or_404(training_id)
    if not t or not allowed_training(claims,t): return jsonify({"message":"Training not found or access denied"}),403
    attendance=attendance_rows(training_id); feedback=feedback_rows(training_id); trainee=trainee_trainer_rows(training_id,t)
    if request.args.get("format")=="xlsx":
        return xlsx_response([("Attendance",attendance),("Feedback",feedback),("Trainee-Trainer",trainee)],f"{t.get('title','training').replace(' ','-')}-report.xlsx")
    feedback_map={r["Employee ID"]:r for r in feedback}; attendance_map={r["Employee ID"]:r for r in attendance}
    combined=[]
    for row in trainee:
        eid=row["Trainee ID"]; combined.append({**row,"Attendance Marked By":attendance_map.get(eid,{}).get("Marked By",""),"Overall Feedback":feedback_map.get(eid,{}).get("Overall",""),"Major Learning":feedback_map.get(eid,{}).get("Major Learning",""),"Remarks":feedback_map.get(eid,{}).get("Remarks",""),"Feedback Submitted At":feedback_map.get(eid,{}).get("Submitted At","")})
    return jsonify({"training":{"id":training_id,"title":t.get("title",""),"trainer":t.get("trainerName",""),"trainingType":t.get("trainingType",""),"status":t.get("status","Upcoming")},"rows":combined,"sheets":{"attendance":attendance,"feedback":feedback,"traineeTrainer":trainee}})

@reports_bp.get("/monthly.xlsx")
@jwt_required()
def monthly_report():
    claims=get_jwt(); year=request.args.get("year",str(datetime.now().year))
    try: year=int(year)
    except ValueError: return jsonify({"message":"Year must be numeric"}),400
    query={}
    if claims.get("role")=="MANAGER":
        ids=[x.get("employeeId") for x in db.employees.find({"department":claims.get("department")},{"employeeId":1,"_id":0}) if x.get("employeeId")]
        query["$or"]=[{"departments":claims.get("department")},{"traineeIds":{"$in":ids}}]
    elif claims.get("role")=="EMPLOYEE": query["$or"]=[{"trainerEmployeeId":claims.get("employeeId")},{"traineeIds":claims.get("employeeId")}]
    trainings=list(db.trainings.find(query).sort("startDate",1))
    details=[]; monthly={m:{"Month":datetime(year,m,1).strftime("%B"),"Year":year,"Scheduled":0,"Completed":0,"Ongoing":0,"Upcoming":0,"Cancelled":0,"One-to-one / Custom":0,"Company-wide / Department":0} for m in range(1,13)}
    for t in trainings:
        raw=t.get("startDate",t.get("trainingDate")); text_date=raw.isoformat() if hasattr(raw,"isoformat") else str(raw or "")
        try: dt=datetime.fromisoformat(text_date.replace("Z","+00:00")); month=dt.month
        except Exception: continue
        if dt.year!=year: continue
        tid=str(t["_id"]); participants=list(db.training_participants.find({"trainingId":tid})); total=len(participants)
        status=t.get("status","Upcoming")
        if status not in {"Completed","Cancelled"}:
            try: status="Ongoing" if datetime.now(dt.tzinfo or None)>=dt else "Upcoming"
            except Exception: status="Upcoming"
        scope="One-to-one / Custom" if total<=1 else "Company-wide / Department"
        a=sum(1 for p in participants if p.get("attendance") in {"Present","Partial"}); f=sum(1 for p in participants if p.get("feedbackStatus")=="Submitted")
        details.append({"Date":text_date,"Training":t.get("title",""),"Type":t.get("trainingType",""),"Trainer":t.get("trainerName",""),"Trainer ID":t.get("trainerEmployeeId",""),"Start":t.get("startDate",t.get("trainingDate")),"End":t.get("endDate",""),"Status":status,"Scope":scope,"Trainees":total,"Attendance Marked":sum(1 for p in participants if p.get("attendance") not in {None,"","Pending"}),"Present / Partial":a,"Feedback Submitted":f,"Attendance %":round(a/total*100,1) if total else 0,"Feedback Completion %":round(f/total*100,1) if total else 0})
        m=monthly[month]; m["Scheduled"]+=1; m[status]=m.get(status,0)+1
        m["One-to-one / Custom" if total<=1 else "Company-wide / Department"]+=1
    summary=list(monthly.values())
    return xlsx_response([("Monthly Summary",summary),("Training Details",details)],f"training-monthly-report-{year}.xlsx")

@reports_bp.get("/training/<training_id>/attendance")
@jwt_required()
def attendance_report(training_id):
    claims=get_jwt(); t=training_or_404(training_id)
    if not t or not allowed_training(claims,t): return jsonify({"message":"Training not found or access denied"}),403
    rows=attendance_rows(training_id)
    if request.args.get("format")=="csv": return csv_response(rows,"attendance-report.csv")
    return jsonify({"training":{"id":training_id,"title":t.get("title","" )},"rows":rows})

@reports_bp.get("/training/<training_id>/feedback")
@jwt_required()
def feedback_report(training_id):
    claims=get_jwt(); t=training_or_404(training_id)
    if not t or not allowed_training(claims,t): return jsonify({"message":"Training not found or access denied"}),403
    rows=feedback_rows(training_id)
    if request.args.get("format")=="csv": return csv_response(rows,"feedback-report.csv")
    return jsonify({"training":{"id":training_id,"title":t.get("title","" )},"rows":rows})

@reports_bp.get("/training/<training_id>/trainee-trainer")
@jwt_required()
def trainee_trainer_report(training_id):
    claims=get_jwt(); t=training_or_404(training_id)
    if not t or not allowed_training(claims,t): return jsonify({"message":"Training not found or access denied"}),403
    rows=trainee_trainer_rows(training_id,t)
    if request.args.get("format")=="csv": return csv_response(rows,"trainee-trainer-report.csv")
    return jsonify({"training":{"id":training_id,"title":t.get("title",""),"trainer":t.get("trainerName","")},"rows":rows})
